"""Slide 6's summary table — a genuine PowerPoint table (GraphicFrame).

Selectable, editable and pasteable. Three header rows: the cable name spanning
its five columns, then the per-cable / ×208 grouping, then the stage. One row
per assembly length. All cell strings come pre-formatted from the JSON dump.
"""

from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from theme import (
    ACCENT,
    BG,
    FONT,
    HAIRLINE,
    INK,
    INK_DIM,
    INK_FAINT,
    MONO,
    PANEL,
    SLIDE_W,
    mix,
    rgb,
)

# "No Style, No Grid" — we paint every cell ourselves.
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

# Schema order of a:tcPr's children; borders must precede the fill.
_TC_ORDER = [
    "a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:lnTlToBr", "a:lnBlToTr", "a:cell3D",
    "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill",
    "a:headers", "a:extLst",
]

# Per-cable column widths (inches): loss, Q̇/cable 40→4, Q̇/cable 4→1, ×208 40→4, ×208 4→1.
# The bundle columns are widest — they are the ones that decide the budget.
COL_W = [1.05, 1.00, 1.00, 1.20, 1.20]
LEN_W = 1.10
NCOL = len(COL_W)


def _clear_style(table):
    tblPr = table._tbl.tblPr
    for e in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(e)
    el = tblPr.makeelement(qn("a:tableStyleId"), {})
    el.text = NO_STYLE
    tblPr.append(el)
    table.first_row = False
    table.horz_banding = False


def _insert_ordered(tcPr, el, tag: str):
    idx = _TC_ORDER.index(tag)
    for child in tcPr:
        ctag = next((t for t in _TC_ORDER if child.tag == qn(t)), None)
        if ctag is not None and _TC_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    tcPr.append(el)


def border(cell, edge: str, color: str, width_pt: float = 1.0):
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[edge]
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn(tag)):
        tcPr.remove(e)
    ln = tcPr.makeelement(
        qn(tag), {"w": str(int(width_pt * 12700)), "cap": "flat", "cmpd": "sng", "algn": "ctr"}
    )
    fill = ln.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color.lstrip("#").upper()}))
    ln.append(fill)
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
    _insert_ordered(tcPr, ln, tag)


def paint(cell, color: str):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(color)


def write(cell, runs, align=PP_ALIGN.RIGHT, pad_x=0.11):
    """runs: list of (text, size, color, bold, spacing, font)."""
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(pad_x)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    for text, size, color, bold, spacing, font in runs:
        r = p.add_run()
        r.text = text
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        if spacing:
            r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def _col(cable_i: int, j: int) -> int:
    return 1 + cable_i * NCOL + j


def summary_table(shapes, data, cables, cable_ids, y):
    n_rows = 3 + len(data["table"])
    n_cols = 1 + 2 * NCOL
    total_w = Inches(LEN_W + 2 * sum(COL_W))
    x = int((SLIDE_W - total_w) / 2)

    gf = shapes.add_table(n_rows, n_cols, x, y, total_w, Inches(4.2))
    tbl = gf.table
    _clear_style(tbl)

    tbl.columns[0].width = Inches(LEN_W)
    for c in range(2):
        for j, w in enumerate(COL_W):
            tbl.columns[_col(c, j)].width = Inches(w)

    tbl.rows[0].height = Inches(0.40)
    tbl.rows[1].height = Inches(0.32)
    tbl.rows[2].height = Inches(0.34)
    for r in range(3, n_rows):
        tbl.rows[r].height = Inches(0.62)

    bundle = {3, 4}
    cold = {2, 4}

    def tint(j: int) -> str:
        return PANEL if j in bundle else BG

    # ── Row 0: cable name, spanning its five columns, in that cable's accent ──
    for r in range(3):
        paint(tbl.cell(r, 0), BG)
    tbl.cell(0, 0).merge(tbl.cell(2, 0))
    write(tbl.cell(0, 0), [("Length", 11.5, INK_DIM, True, 0.4, FONT)], align=PP_ALIGN.LEFT)

    for c, cid in enumerate(cable_ids):
        accent = ACCENT[cid]
        for j in range(NCOL):
            paint(tbl.cell(0, _col(c, j)), BG)
        first, last = _col(c, 0), _col(c, NCOL - 1)
        tbl.cell(0, first).merge(tbl.cell(0, last))
        write(tbl.cell(0, first), [(cables[cid]["name"], 12.5, accent, True, 0.2, FONT)],
              align=PP_ALIGN.CENTER)
        border(tbl.cell(0, first), "B", mix(BG, accent, 0.45), 1.5)

    # ── Row 1: per-cable vs. bundle. The harness is labelled right in the header. ──
    groups = [(0, 1, "Loss"), (1, 2, "Q̇ per cable"), (3, 2, data["meta"]["harnessLabel"])]
    for c in range(2):
        for j0, span, label in groups:
            for j in range(j0, j0 + span):
                paint(tbl.cell(1, _col(c, j)), tint(j))
            head = tbl.cell(1, _col(c, j0))
            if span > 1:
                head.merge(tbl.cell(1, _col(c, j0 + span - 1)))
            size = 8.0 if j0 == 3 else 8.4
            # No letter-spacing: `spc` breaks the Q̇ combining-mark cluster.
            write(head, [(label.upper(), size, INK_FAINT, False, 0, MONO)],
                  align=PP_ALIGN.CENTER)

    # ── Row 2: the stage each column reports ──
    subs = ["2 GHz", "40→4 K", "4→1 K", "40→4 K", "4→1 K"]
    for c in range(2):
        for j, sub in enumerate(subs):
            cell = tbl.cell(2, _col(c, j))
            paint(cell, tint(j))
            color = INK_FAINT if j in cold else INK_DIM
            write(cell, [(sub, 8.6, color, j not in cold, 0.3, MONO)])
            border(cell, "B", HAIRLINE, 1.0)
    border(tbl.cell(2, 0), "B", HAIRLINE, 1.0)

    # ── Data rows ──
    for r, row in enumerate(data["table"]):
        rr = r + 3
        paint(tbl.cell(rr, 0), BG)
        write(tbl.cell(rr, 0), [(row["length"], 12.5, INK, True, 0, FONT)], align=PP_ALIGN.LEFT)
        if r:
            border(tbl.cell(rr, 0), "T", HAIRLINE, 0.75)

        for c in range(2):
            for j in range(NCOL):
                cell = tbl.cell(rr, _col(c, j))
                d = row["cols"][c * NCOL + j]
                paint(cell, tint(j))

                is_bundle = j in bundle
                is_cold = j in cold
                if is_cold:
                    size, color, bold = 11.0, INK_FAINT, False
                elif is_bundle:
                    size, color, bold = 12.5, INK, True
                else:
                    size, color, bold = 11.5, INK_DIM, False

                write(cell, [
                    (d["v"], size, color, bold, 0, MONO),
                    (" " + d["u"], size * 0.72, INK_FAINT, False, 0, MONO),
                ])
                if r:
                    border(cell, "T", HAIRLINE, 0.75)

    return gf
